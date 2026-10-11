"""
AgriOne Presentation Generator
Builds a 36-slide widescreen (16:9) PowerPoint presentation for AgriOne
Team: SILENT CODERS
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Design System Colors
C_BG_DARK = RGBColor(8, 20, 14)       # Deep obsidian emerald
C_CARD_BG = RGBColor(14, 34, 25)      # Dark emerald container
C_CARD_ALT = RGBColor(18, 44, 32)     # Alternate elevated card
C_BORDER = RGBColor(27, 67, 50)       # Border tone
C_ACCENT_GREEN = RGBColor(16, 185, 129) # Vivid emerald
C_ACCENT_CYAN = RGBColor(6, 182, 212)   # Cyber cyan
C_ACCENT_LIME = RGBColor(132, 204, 22)  # Lime accent
C_TEXT_WHITE = RGBColor(248, 250, 252)  # Slate white
C_TEXT_MUTED = RGBColor(148, 163, 184)  # Slate muted gray
C_TEXT_DIM = RGBColor(100, 116, 139)    # Dim gray
C_CARD_ALERT = RGBColor(45, 20, 20)    # Warning card bg
C_BORDER_ALERT = RGBColor(180, 50, 50)  # Warning border
C_ACCENT_RED = RGBColor(239, 68, 68)    # Alert red
C_HIGHLIGHT = RGBColor(52, 211, 153)   # Bright highlight

FONT_TITLE = "Segoe UI"
FONT_BODY = "Segoe UI"

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide
    
    workspace_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(workspace_dir, "assets")
    sc_logo_path = os.path.join(assets_dir, "silent_coders_logo.jpg")
    ag_logo_path = os.path.join(assets_dir, "agrione_logo.jpg")
    ag_hero_path = os.path.join(assets_dir, "agrione_hero.jpg")

    def add_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, slide_num, category, title, subtitle=None):
        # Top banner line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.04))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_BORDER
        top_bar.line.fill.background()

        # Category & Brand tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.48), Inches(9.0), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = f"AGRIONE • {category.upper()}"
        p_cat.font.name = FONT_TITLE
        p_cat.font.size = Pt(9.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_ACCENT_GREEN

        # Slide Number tag
        num_box = slide.shapes.add_textbox(Inches(10.5), Inches(0.48), Inches(2.033), Inches(0.3))
        tf_num = num_box.text_frame
        tf_num.margin_left = tf_num.margin_right = tf_num.margin_top = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.text = f"SLIDE {slide_num:02d} / 36"
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.name = FONT_TITLE
        p_num.font.size = Pt(9.5)
        p_num.font.color.rgb = C_TEXT_DIM

        # Slide Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.733), Inches(0.55))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_WHITE

        # Subtitle if present
        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.35))
            tf_sub = sub_box.text_frame
            tf_sub.word_wrap = True
            tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
            p_sub = tf_sub.paragraphs[0]
            p_sub.text = subtitle
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(12)
            p_sub.font.color.rgb = C_ACCENT_CYAN

    def add_card(slide, left, top, width, height, title=None, body_items=None, border_color=C_BORDER, bg_color=C_CARD_BG, title_color=C_ACCENT_GREEN):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)

        padding = Inches(0.18)
        content_box = slide.shapes.add_textbox(left + padding, top + padding, width - (padding * 2), height - (padding * 2))
        tf = content_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0

        p_idx = 0
        if title:
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.font.name = FONT_TITLE
            p_title.font.size = Pt(13)
            p_title.font.bold = True
            p_title.font.color.rgb = title_color
            p_title.space_after = Pt(6)
            p_idx += 1

        if body_items:
            for item in body_items:
                p = tf.add_paragraph() if (title or p_idx > 0) else tf.paragraphs[0]
                p.text = f"• {item}" if not item.startswith("—") and not item.startswith("▶") else item
                p.font.name = FONT_BODY
                p.font.size = Pt(10.5)
                p.font.color.rgb = C_TEXT_WHITE
                p.space_after = Pt(4)
                p_idx += 1

        return card

    # =========================================================================
    # SLIDE 1: TEAM INTRODUCTION
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_slide_background(s1)

    # Top brand badge
    top_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.6), Inches(4.2), Inches(0.42))
    top_badge.fill.solid()
    top_badge.fill.fore_color.rgb = C_CARD_BG
    top_badge.line.color.rgb = C_ACCENT_GREEN
    top_badge.line.width = Pt(1)
    tf_tb = top_badge.text_frame
    tf_tb.paragraphs[0].text = "HACKATHON / COLLEGE PROJECT PRESENTATION"
    tf_tb.paragraphs[0].font.name = FONT_TITLE
    tf_tb.paragraphs[0].font.size = Pt(9.5)
    tf_tb.paragraphs[0].font.bold = True
    tf_tb.paragraphs[0].font.color.rgb = C_ACCENT_GREEN
    tf_tb.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Title & Subtitle
    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.2), Inches(7.5), Inches(1.6))
    tf1 = tb_title.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "SILENT CODERS"
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "AgriOne – Intelligent Digital Agriculture Platform"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(19)
    p2.font.bold = True
    p2.font.color.rgb = C_ACCENT_CYAN
    p2.space_before = Pt(6)

    p3 = tf1.add_paragraph()
    p3.text = "AI-Driven Ecosystem for Real-Time Agricultural Decision Support"
    p3.font.name = FONT_BODY
    p3.font.size = Pt(12)
    p3.font.color.rgb = C_TEXT_MUTED
    p3.space_before = Pt(4)

    # Team Logo on Right
    if os.path.exists(sc_logo_path):
        s1.shapes.add_picture(sc_logo_path, Inches(8.8), Inches(0.9), Inches(3.6), Inches(3.6))

    # Team Cards Section
    # Team Leader Card (Highlighted)
    leader_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.4), Inches(11.733), Inches(1.5))
    leader_card.fill.solid()
    leader_card.fill.fore_color.rgb = C_CARD_ALT
    leader_card.line.color.rgb = C_ACCENT_GREEN
    leader_card.line.width = Pt(1.5)

    tf_lead = leader_card.text_frame
    tf_lead.word_wrap = True
    p_ltag = tf_lead.paragraphs[0]
    p_ltag.text = "★ TEAM LEADER"
    p_ltag.font.name = FONT_TITLE
    p_ltag.font.size = Pt(11)
    p_ltag.font.bold = True
    p_ltag.font.color.rgb = C_ACCENT_GREEN

    p_lname = tf_lead.add_paragraph()
    p_lname.text = "Iyaan Nagari"
    p_lname.font.name = FONT_TITLE
    p_lname.font.size = Pt(22)
    p_lname.font.bold = True
    p_lname.font.color.rgb = C_TEXT_WHITE
    p_lname.space_before = Pt(2)

    p_lrole = tf_lead.add_paragraph()
    p_lrole.text = "Team Lead • System Architecture & Project Direction"
    p_lrole.font.name = FONT_BODY
    p_lrole.font.size = Pt(11)
    p_lrole.font.color.rgb = C_ACCENT_CYAN
    p_lrole.space_before = Pt(2)

    # 3 Member Cards
    members = [
        ("Bhuvan Kadakabhavi", "Team Member • Silent Coders"),
        ("Jafar Momin", "Team Member • Silent Coders"),
        ("Aiman Nadaf", "Team Member • Silent Coders")
    ]
    card_w = Inches(3.75)
    gap = Inches(0.24)
    for idx, (m_name, m_sub) in enumerate(members):
        m_left = Inches(0.8) + idx * (card_w + gap)
        m_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_left, Inches(5.15), card_w, Inches(1.65))
        m_card.fill.solid()
        m_card.fill.fore_color.rgb = C_CARD_BG
        m_card.line.color.rgb = C_BORDER
        m_card.line.width = Pt(1)

        tf_m = m_card.text_frame
        tf_m.word_wrap = True
        p_mtag = tf_m.paragraphs[0]
        p_mtag.text = f"TEAM MEMBER 0{idx+1}"
        p_mtag.font.name = FONT_TITLE
        p_mtag.font.size = Pt(9.5)
        p_mtag.font.bold = True
        p_mtag.font.color.rgb = C_ACCENT_CYAN

        p_mname = tf_m.add_paragraph()
        p_mname.text = m_name
        p_mname.font.name = FONT_TITLE
        p_mname.font.size = Pt(16)
        p_mname.font.bold = True
        p_mname.font.color.rgb = C_TEXT_WHITE
        p_mname.space_before = Pt(4)

        p_msub = tf_m.add_paragraph()
        p_msub.text = m_sub
        p_msub.font.name = FONT_BODY
        p_msub.font.size = Pt(10)
        p_msub.font.color.rgb = C_TEXT_MUTED
        p_msub.space_before = Pt(3)

    # =========================================================================
    # SLIDE 2: PROJECT INTRODUCTION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_background(s2)
    add_header(s2, 2, "Introduction", "AgriOne – Intelligent Digital Agriculture Platform", "“AI-Powered Intelligence for Smarter Farming.”")

    # Left Column: Overview & Philosophy
    left_w = Inches(6.8)
    add_card(s2, Inches(0.8), Inches(1.8), left_w, Inches(2.3),
             title="The Core Mission of AgriOne",
             body_items=[
                 "AgriOne is an intelligent digital agriculture platform connecting farmers with cutting-edge AI and real-time agricultural intelligence.",
                 "Combines crop health analysis, weather patterns, soil dynamics, mandi market intelligence, and farm management tools.",
                 "Transforms raw farm telemetry into simple, actionable daily decisions within one unified ecosystem."
             ],
             border_color=C_BORDER, bg_color=C_CARD_BG, title_color=C_HIGHLIGHT)

    # Value Pipeline on Left
    pipe_items = [
        ("Farmer", "Ground reality, observations, crops"),
        ("Data", "Soil, weather, satellites, mandi prices"),
        ("AI + Rule Engine", "Neural models & validated agronomy"),
        ("Insights", "Precise forecasts & health diagnoses"),
        ("Better Decisions", "Maximized yield & reduced risk")
    ]
    pipe_top = Inches(4.3)
    pipe_w = Inches(6.8)
    pipe_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), pipe_top, pipe_w, Inches(2.6))
    pipe_card.fill.solid()
    pipe_card.fill.fore_color.rgb = C_CARD_ALT
    pipe_card.line.color.rgb = C_ACCENT_GREEN
    pipe_card.line.width = Pt(1)

    tf_pipe = pipe_card.text_frame
    tf_pipe.word_wrap = True
    p_ptit = tf_pipe.paragraphs[0]
    p_ptit.text = "CORE VALUE CHAIN FLOW"
    p_ptit.font.name = FONT_TITLE
    p_ptit.font.size = Pt(11)
    p_ptit.font.bold = True
    p_ptit.font.color.rgb = C_ACCENT_GREEN

    p_flow = tf_pipe.add_paragraph()
    p_flow.text = "Farmer  ➔  Data  ➔  AI + Rule Engine  ➔  Insights  ➔  Better Decisions"
    p_flow.font.name = FONT_TITLE
    p_flow.font.size = Pt(13)
    p_flow.font.bold = True
    p_flow.font.color.rgb = C_TEXT_WHITE
    p_flow.space_before = Pt(6)
    p_flow.space_after = Pt(6)

    for step, desc in pipe_items:
        p_step = tf_pipe.add_paragraph()
        p_step.text = f"• {step}: {desc}"
        p_step.font.name = FONT_BODY
        p_step.font.size = Pt(9.5)
        p_step.font.color.rgb = C_TEXT_MUTED

    # Right Column: Hero Visual or Graphic
    if os.path.exists(ag_hero_path):
        s2.shapes.add_picture(ag_hero_path, Inches(7.8), Inches(1.8), Inches(4.733), Inches(5.1))
    else:
        add_card(s2, Inches(7.8), Inches(1.8), Inches(4.733), Inches(5.1),
                 title="Next-Generation Smart Agriculture",
                 body_items=["Unified digital command center", "Real-time edge intelligence", "Accessible to every farmer"])

    # =========================================================================
    # SLIDE 3: THE PROBLEM
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_background(s3)
    add_header(s3, 3, "Problem Statement", "Problems Faced by Modern Farmers", "Major hurdles that cause financial loss, yield reduction, and operational uncertainty.")

    problems = [
        ("Unpredictable Weather", "Sudden climate variations destroy unharvested crops without early localized warnings."),
        ("Crop Disease Identification", "Delayed or incorrect pest diagnosis leads to rampant crop failure and toxic chemical misuse."),
        ("Improper Irrigation Decisions", "Over-watering or underwatering degrades soil health and wastes scarce water resources."),
        ("Soil Health Uncertainty", "Lack of accessible soil testing leaves farmers unaware of critical N-P-K nutrient deficits."),
        ("Fragmented Information", "Farming data is scattered across isolated apps, TV broadcasts, and hearsay."),
        ("Mandi Price Opacity", "Middlemen exploit lack of real-time market transparency, slashing farmer profits."),
        ("Supply-Chain Bottlenecks", "Inefficient transport logistics cause massive post-harvest spoilage."),
        ("Limited Expert Access", "Agronomists and agricultural university extension centers are distant and understaffed."),
        ("Language & Tech Barriers", "Complex English software alienates rural and regional language speakers."),
        ("Manual Farm Management", "Paper-based tracking results in forgotten schedules, missed spray cycles, and disorganized records.")
    ]

    col_w = Inches(2.22)
    card_h = Inches(2.35)
    gap_x = Inches(0.15)
    row_gap = Inches(0.2)

    for i, (p_title, p_desc) in enumerate(problems):
        row = i // 5
        col = i % 5
        c_left = Inches(0.8) + col * (col_w + gap_x)
        c_top = Inches(1.85) + row * (card_h + row_gap)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, col_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_ALERT
        card.line.color.rgb = C_BORDER_ALERT
        card.line.width = Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.12)
        
        p_badge = tf.paragraphs[0]
        p_badge.text = f"ISSUE 0{i+1}" if i < 9 else f"ISSUE {i+1}"
        p_badge.font.name = FONT_TITLE
        p_badge.font.size = Pt(8.5)
        p_badge.font.bold = True
        p_badge.font.color.rgb = C_ACCENT_RED

        p_t = tf.add_paragraph()
        p_t.text = p_title
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = C_TEXT_WHITE
        p_t.space_before = Pt(3)

        p_d = tf.add_paragraph()
        p_d.text = p_desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(9)
        p_d.font.color.rgb = C_TEXT_MUTED
        p_d.space_before = Pt(3)

    # =========================================================================
    # SLIDE 4: OUR SOLUTION
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_background(s4)
    add_header(s4, 4, "Solution Architecture", "What is AgriOne?", "A unified intelligent ecosystem bridging scientific agronomy with practical farm decision-making.")

    # Core Philosophy Box
    phil_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(1.0))
    phil_box.fill.solid()
    phil_box.fill.fore_color.rgb = C_CARD_ALT
    phil_box.line.color.rgb = C_ACCENT_GREEN
    phil_box.line.width = Pt(1.5)
    tf_phil = phil_box.text_frame
    tf_phil.word_wrap = True
    p_ph1 = tf_phil.paragraphs[0]
    p_ph1.text = "CORE ARCHITECTURAL PHILOSOPHY"
    p_ph1.font.name = FONT_TITLE
    p_ph1.font.size = Pt(10)
    p_ph1.font.bold = True
    p_ph1.font.color.rgb = C_ACCENT_GREEN

    p_ph2 = tf_phil.add_paragraph()
    p_ph2.text = "“AI interprets ➔ Rule Engine validates & calculates ➔ AgriOne provides actionable decisions”"
    p_ph2.font.name = FONT_TITLE
    p_ph2.font.size = Pt(15)
    p_ph2.font.bold = True
    p_ph2.font.color.rgb = C_TEXT_WHITE
    p_ph2.space_before = Pt(3)

    # 3 Workflow Pillars
    pillars = [
        ("01. AI Interprets", C_ACCENT_CYAN, [
            "Natural language query comprehension",
            "Multi-modal crop disease image scanning",
            "Contextual conversational understanding",
            "Extracts intent, location, crop & soil nuances"
        ]),
        ("02. Rule Engine Validates & Calculates", C_ACCENT_GREEN, [
            "Mathematical agronomy equations",
            "Threshold validation for moisture & weather",
            "Fertilizer N-P-K requirement calculation",
            "Strict agricultural safety boundary enforcement"
        ]),
        ("03. Actionable Decisions Delivered", C_HIGHLIGHT, [
            "Instant farmer advisory in native languages",
            "Precision irrigation timer recommendations",
            "Tailored organic & chemical spray guidance",
            "High-profit market selling time suggestions"
        ])
    ]

    p_w = Inches(3.75)
    p_gap = Inches(0.24)
    for idx, (p_name, p_color, p_bullets) in enumerate(pillars):
        p_left = Inches(0.8) + idx * (p_w + p_gap)
        add_card(s4, p_left, Inches(3.05), p_w, Inches(3.85),
                 title=p_name,
                 body_items=p_bullets,
                 border_color=p_color,
                 bg_color=C_CARD_BG,
                 title_color=p_color)

    # =========================================================================
    # SLIDE 5: COMPLETE AGRIONE ECOSYSTEM
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_background(s5)
    add_header(s5, 5, "Ecosystem Overview", "Complete AgriOne Ecosystem", "A central intelligence hub connecting 18 specialized agricultural pillars.")

    modules = [
        "AI Farmer Assistant", "Crop Health & Disease", "Smart Irrigation",
        "Weather Intelligence", "Crop Intelligence", "Soil Health",
        "Crop Recommendation", "Mandi / Market Prices", "Farm Management",
        "Supply Chain", "Marketplace", "Livestock & Dairy",
        "Satellite & Drone Intel", "Sustainable Farming", "Finance & Insurance",
        "Analytics & Trends", "AI Voice Interaction", "Multilingual Platform"
    ]

    mod_w = Inches(3.75)
    mod_h = Inches(0.7)
    gap_x = Inches(0.24)
    gap_y = Inches(0.14)

    for i, mod in enumerate(modules):
        r = i // 3
        c = i % 3
        m_x = Inches(0.8) + c * (mod_w + gap_x)
        m_y = Inches(1.8) + r * (mod_h + gap_y)

        btn = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, m_x, m_y, mod_w, mod_h)
        btn.fill.solid()
        btn.fill.fore_color.rgb = C_CARD_BG
        btn.line.color.rgb = C_ACCENT_GREEN if i % 2 == 0 else C_ACCENT_CYAN
        btn.line.width = Pt(1)

        tf = btn.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        p = tf.paragraphs[0]
        p.text = f"#{i+1:02d}  {mod}"
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_WHITE

    # =========================================================================
    # SLIDE 6: FARMER DASHBOARD
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_background(s6)
    add_header(s6, 6, "Farmer Dashboard", "Personalized Farm Command Center", "One consolidated real-time screen tracking every vital parameter of the farm.")

    dash_widgets = [
        ("Farm Overview & Acreage", ["Boundaries: 4.8 Hectares", "Active parcels: 3 zones", "Current status: Active Growth"]),
        ("Location & GPS Setup", ["Geo-tagged: Dharwad, Karnataka", "Coordinates: 15.45° N, 75.00° E", "Elevation & Agro-zone calibrated"]),
        ("Crop & Variety Profile", ["Primary: Soybeans (JS-335)", "Secondary: Wheat (HD-2967)", "Stage: Vegetative (Day 42)"]),
        ("Soil & Nutrition Index", ["Soil Type: Medium Black Soil", "Moisture: 42% (Optimal)", "NPK: N-Medium, P-High, K-Adequate"]),
        ("Smart Irrigation Status", ["Next Cycle: 06:00 AM Tomorrow", "Water Volume: 12,000 Liters", "Mode: Automated Drip Calculation"]),
        ("Weather & Microclimate", ["Temp: 29°C | Humidity: 68%", "Precipitation probability: 15%", "Wind: 9 km/h North-East"]),
        ("Crop Health Scanner", ["Health Index: 94% Healthy", "Recent scans: 2 images analyzed", "Alert: Mild leaf spot detected"]),
        ("Daily Farm Operations", ["Tasks Pending: 2 Field Activities", "Fertigation scheduled in 48h", "Equipment ready: Tractor & Sprayer"])
    ]

    dw_w = Inches(2.78)
    dw_h = Inches(2.4)
    gap_x = Inches(0.18)
    gap_y = Inches(0.2)

    for i, (dw_tit, dw_items) in enumerate(dash_widgets):
        r = i // 4
        c = i % 4
        d_x = Inches(0.8) + c * (dw_w + gap_x)
        d_y = Inches(1.85) + r * (dw_h + gap_y)

        add_card(s6, d_x, d_y, dw_w, dw_h,
                 title=dw_tit,
                 body_items=dw_items,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 7: SMART LOCATION & FARM SETUP
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_background(s7)
    add_header(s7, 7, "Location Awareness", "Location-Aware Agriculture", "Hyper-local parameters drive precision insights down to the farmer's specific field.")

    # Flow Banner
    f_bar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.8))
    f_bar.fill.solid()
    f_bar.fill.fore_color.rgb = C_CARD_ALT
    f_bar.line.color.rgb = C_ACCENT_CYAN
    f_bar.line.width = Pt(1)
    tf_f = f_bar.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.text = "CORE FLOW:  Location  ➔  Farm  ➔  Weather  ➔  Crop  ➔  Recommendation"
    p_f.font.name = FONT_TITLE
    p_f.font.size = Pt(13)
    p_f.font.bold = True
    p_f.font.color.rgb = C_ACCENT_CYAN
    p_f.alignment = PP_ALIGN.CENTER

    loc_features = [
        ("1. Farmer Selects Location", [
            "Seamless village, taluk, and district search",
            "Instant GPS pinpointing with one tap",
            "Geo-boundary polygon mapping on satellite view"
        ]),
        ("2. Agro-Climatic Intelligence", [
            "Auto-fetches agro-climatic sub-zone data",
            "Integrates local historical monsoon records",
            "Retrieves state and regional government mandates"
        ]),
        ("3. Precision Farm Profiling", [
            "Binds soil characteristics to elevation data",
            "Tailors weather forecast to local field radius",
            "Enables personalized crop recommendations"
        ])
    ]

    for idx, (ltit, litems) in enumerate(loc_features):
        lx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s7, lx, Inches(2.85), Inches(3.75), Inches(4.1),
                 title=ltit,
                 body_items=litems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 8: AI FARMER ASSISTANT
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_background(s8)
    add_header(s8, 8, "AI Assistant", "Your AI Agricultural Assistant", "Intelligent agronomic advice answering the actual questions asked by the farmer.")

    # Left: Real query examples
    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Real Farmer Queries Handled",
             body_items=[
                 "“Why are my tomato leaves curling and developing brown spots?”",
                 "“Should I irrigate my wheat field tomorrow given the rain forecast?”",
                 "“What is the recommended NPK fertilizer dosage for black soil cotton?”",
                 "“Which mandi is offering the highest price for soybean within 50 km?”",
                 "“How can I manage stem borers organically without costly chemicals?”",
                 "— Answers actual user questions with domain rigor, not generic canned responses."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_GREEN)

    # Right: Capabilities
    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="Assistant Capabilities & Guardrails",
             body_items=[
                 "Multi-Domain Knowledge: Agronomy, botany, meteorology & economics.",
                 "Context-Aware Memory: Knows the farmer's soil, crop stage, and location.",
                 "Rule-Engine Grounding: Cross-verifies chemical dosage with safe limits.",
                 "Actionable Guidance: Step-by-step instructions with exact timings.",
                 "Follow-Up Continuity: Retains conversational context across sessions.",
                 "Multi-modal: Understands both voice queries and uploaded leaf photos."
             ],
             border_color=C_ACCENT_CYAN,
             bg_color=C_CARD_ALT,
             title_color=C_ACCENT_CYAN)

    # =========================================================================
    # SLIDE 9: AI VOICE INTERACTION
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_background(s9)
    add_header(s9, 9, "Voice Intelligence", "Talk to AgriOne", "Enabling frictionless hands-free farming advisory right in the muddy field.")

    # Flow Banner
    v_flow = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.85))
    v_flow.fill.solid()
    v_flow.fill.fore_color.rgb = C_CARD_ALT
    v_flow.line.color.rgb = C_ACCENT_CYAN
    v_flow.line.width = Pt(1.5)
    tf_vf = v_flow.text_frame
    p_vf1 = tf_vf.paragraphs[0]
    p_vf1.text = "END-TO-END VOICE PIPELINE"
    p_vf1.font.name = FONT_TITLE
    p_vf1.font.size = Pt(9.5)
    p_vf1.font.bold = True
    p_vf1.font.color.rgb = C_ACCENT_CYAN
    p_vf2 = tf_vf.add_paragraph()
    p_vf2.text = "Voice Input  ➔  AI Understanding  ➔  Agronomic Processing  ➔  Voice Response"
    p_vf2.font.name = FONT_TITLE
    p_vf2.font.size = Pt(13)
    p_vf2.font.bold = True
    p_vf2.font.color.rgb = C_TEXT_WHITE
    p_vf2.space_before = Pt(3)

    v_cards = [
        ("Speech Recognition (ASR)", [
            "Trained for rural dialects & regional accents",
            "Filters ambient background tractor & wind noise",
            "High accuracy on regional agricultural terms"
        ]),
        ("Intent & Agronomy Engine", [
            "Decodes phonetic spelling of crop & pesticide names",
            "Executes soil & weather validation logic",
            "Prepares concise, easily listenable answers"
        ]),
        ("Natural Text-to-Speech (TTS)", [
            "Human-like regional voice synthesis",
            "Speaks in farmer's chosen native language",
            "Completely hands-free operation in the field"
        ])
    ]

    for idx, (vtit, vitems) in enumerate(v_cards):
        vx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s9, vx, Inches(2.95), Inches(3.75), Inches(4.0),
                 title=vtit,
                 body_items=vitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 10: MULTILINGUAL PLATFORM
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_background(s10)
    add_header(s10, 10, "Accessibility", "Agriculture Without Language Barriers", "“One Platform. Multiple Languages. Better Accessibility.”")

    # Banner
    m_ban = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(0.9))
    m_ban.fill.solid()
    m_ban.fill.fore_color.rgb = C_CARD_ALT
    m_ban.line.color.rgb = C_ACCENT_GREEN
    m_ban.line.width = Pt(1)
    tf_mb = m_ban.text_frame
    p_mb = tf_mb.paragraphs[0]
    p_mb.text = "When a farmer switches language, the entire ecosystem adapts synchronously:"
    p_mb.font.name = FONT_TITLE
    p_mb.font.size = Pt(13)
    p_mb.font.bold = True
    p_mb.font.color.rgb = C_TEXT_WHITE
    p_mb.alignment = PP_ALIGN.CENTER

    m_features = [
        ("Full UI & Navigation", [
            "Dashboard labels & metrics translate dynamically",
            "Action buttons & form inputs switch instantly",
            "Culturally intuitive iconography for effortless navigation"
        ]),
        ("Conversational AI Support", [
            "AI assistant reads & responds in the selected tongue",
            "Preserves local colloquial farming expressions",
            "Explains complex agricultural advice in simple words"
        ]),
        ("Voice In & Voice Out", [
            "Speaks back to the farmer in their chosen dialect",
            "Supports low-literacy farmers through audio-first design",
            "Breaks the digital divide across rural communities"
        ])
    ]

    for idx, (mtit, mitems) in enumerate(m_features):
        mx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s10, mx, Inches(3.0), Inches(3.75), Inches(3.95),
                 title=mtit,
                 body_items=mitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 11: AI CROP HEALTH & DISEASE DETECTION
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_slide_background(s11)
    add_header(s11, 11, "Computer Vision", "AI Crop Health Scanner", "Instant disease diagnosis and treatment recommendations from real camera photos.")

    steps_cv = [
        ("1. Image Upload", "Farmer uploads an actual leaf/crop photograph from field camera."),
        ("2. AI Analysis", "Deep neural network evaluates leaf lesions, discoloration, and patterns."),
        ("3. Identification", "Identifies exact disease pathogen (fungal, bacterial, viral, or pest)."),
        ("4. Assessment", "Determines disease severity percentage and infection stage."),
        ("5. Possible Cause", "Pinpoints underlying causes: humidity spikes, soil pathogens, or nutrient lack."),
        ("6. Action Plan", "Generates chemical/organic remedies, dosage per liter, and spray timetable.")
    ]

    for idx, (s_name, s_desc) in enumerate(steps_cv):
        row = idx // 3
        col = idx % 3
        sx = Inches(0.8) + col * (Inches(3.75) + Inches(0.24))
        sy = Inches(1.85) + row * (Inches(2.5) + Inches(0.2))

        add_card(s11, sx, sy, Inches(3.75), Inches(2.4),
                 title=s_name,
                 body_items=[s_desc],
                 border_color=C_ACCENT_GREEN if idx in [1, 5] else C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_ACCENT_GREEN if idx in [1, 5] else C_TEXT_WHITE)

    # =========================================================================
    # SLIDE 12: SMART IRRIGATION
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_slide_background(s12)
    add_header(s12, 12, "Resource Management", "Intelligent Water Management", "Crop + Soil + Weather + Moisture ➔ AI / Rule Engine ➔ Irrigation Recommendation")

    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Inputs Analyzed for Irrigation",
             body_items=[
                 "Crop Phenology: Current growth stage (e.g. flowering requires higher water, maturity requires less).",
                 "Soil Water Capacity: Sandy vs clay vs black soil infiltration rates.",
                 "Weather Forecast: Postpones irrigation if rainfall > 10mm is forecast within 24h.",
                 "Moisture Telemetry: Real-time sensor readings or evapotranspiration calculations.",
                 "Prevents root rot from over-watering while protecting against moisture stress."
             ],
             border_color=C_ACCENT_CYAN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_CYAN)

    add_card(s12, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="Operational Highlights",
             body_items=[
                 "Efficient Water Conservation: Saves 30% to 40% of agricultural water.",
                 "Smart Scheduling: Recommends early morning/evening irrigation to reduce evaporation loss.",
                 "Pump Energy Optimization: Minimizes electricity consumption during peak tariff hours.",
                 "Actionable Schedule: 'Run Drip Line Zone A for 45 minutes at 06:00 AM tomorrow'.",
                 "Automated Alerts: Push notifications sent when soil hits critical wilting threshold."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_ALT,
             title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 13: WEATHER & CROP INTELLIGENCE
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_slide_background(s13)
    add_header(s13, 13, "Microclimate Insights", "From Weather Data to Farming Decisions", "Weather Data ➔ AI Interpretation ➔ Concrete Farmer Action")

    w_metrics = [
        ("Temperature & Heat Index", "Monitors heatwaves; advises foliar anti-transpirant sprays during extreme thermal stress."),
        ("Precipitation & Rain Probability", "7-day localized rainfall forecast prevents fertilizer wash-off and saves irrigation costs."),
        ("Relative Humidity & Dew Point", "Elevated humidity triggers early warnings for high-risk fungal blight outbreaks."),
        ("Wind Velocity & Direction", "Advises whether aerial or tractor spray operations can proceed without spray drift.")
    ]

    for idx, (wtit, wdesc) in enumerate(w_metrics):
        wx = Inches(0.8) + (idx % 2) * (Inches(5.7) + Inches(0.333))
        wy = Inches(1.85) + (idx // 2) * (Inches(2.5) + Inches(0.2))

        add_card(s13, wx, wy, Inches(5.7), Inches(2.45),
                 title=wtit,
                 body_items=[
                     wdesc,
                     "AgriOne Translation: Converts raw mm/kmh numbers into exact farmer action instructions."
                 ],
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 14: SOIL HEALTH
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_slide_background(s14)
    add_header(s14, 14, "Agronomy & Nutrients", "Understand Your Soil", "Soil Data ➔ Analysis ➔ Crop Suitability ➔ Precision Nutrient Recommendation")

    s_pillars = [
        ("Soil Characteristics", [
            "Soil Texture & Type (Black, Red, Loamy, Sandy)",
            "pH Level & Electrical Conductivity (EC)",
            "Organic Carbon (OC) percentage"
        ]),
        ("Nutrient Balances", [
            "Primary Nutrients: Nitrogen (N), Phosphorus (P), Potassium (K)",
            "Secondary & Micronutrients: Zinc, Iron, Sulfur, Boron",
            "Deficiency detection vs crop demand"
        ]),
        ("Actionable Output", [
            "Custom fertilizer blend calculation per acre",
            "Organic amendment recommendations (Bio-fertilizer, FYM)",
            "Soil conditioning advice to restore fertility long-term"
        ])
    ]

    for idx, (stit, sitems) in enumerate(s_pillars):
        sx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s14, sx, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=stit,
                 body_items=sitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 15: CROP RECOMMENDATION
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_slide_background(s15)
    add_header(s15, 15, "Crop Planning", "What Should I Grow?", "Scientific matching of seasonal conditions, agro-climatic factors, and profitability.")

    add_card(s15, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Multi-Factor Input Analysis",
             body_items=[
                 "Agro-climatic Zone: Altitude, latitude, and regional rainfall patterns.",
                 "Soil Profile Compatibility: pH, drainage, depth, and NPK indices.",
                 "Seasonal Window: Kharif, Rabi, or Zaid crop cycles.",
                 "Water Availability: Rainfed vs canal vs borewell drip capacity.",
                 "Historical Mandi Trends: Projected demand and price realization."
             ],
             border_color=C_ACCENT_CYAN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_CYAN)

    add_card(s15, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="Structured Recommendation Output",
             body_items=[
                 "Recommended Crop: Specific crop variety tailored to farm profile.",
                 "Agronomic Rationale: Why this crop excels under current conditions.",
                 "Supporting Conditions: Required seed rate, spacing, and irrigation.",
                 "Risk Assessment: Vulnerability to pests and climate extremes.",
                 "Expected Yield Range: Realistic harvest output estimate."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_ALT,
             title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 16: MANDI & MARKET INTELLIGENCE
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    add_slide_background(s16)
    add_header(s16, 16, "Market Transparency", "Know the Market Before You Sell", "Crop ➔ Market Data ➔ Price Intelligence ➔ Better Selling Decision")

    mandi_cols = [
        ("Live Mandi Feeds", [
            "Real-time prices from nearby APMC mandis",
            "Min, Modal, and Max price per quintal",
            "Distance-adjusted net realization"
        ]),
        ("Price Trend Analysis", [
            "Weekly & monthly market price trajectory",
            "Arrival volume vs price elasticity",
            "Predictive seasonality indicators"
        ]),
        ("Optimal Selling Timing", [
            "Sell now vs store advisory based on trends",
            "Comparison across 5 neighboring markets",
            "Empowers farmers against middleman price suppression"
        ])
    ]

    for idx, (mtit, mitems) in enumerate(mandi_cols):
        mx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s16, mx, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=mtit,
                 body_items=mitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 17: FARM MANAGEMENT
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    add_slide_background(s17)
    add_header(s17, 17, "Operations", "Digital Farm Management", "Eliminating paper records with a streamlined digital agricultural ledger.")

    fm_features = [
        ("Farm & Parcel Profile", "Digitizes land boundaries, acreage, soil types, and water infrastructure across all plots."),
        ("Crop Record Ledger", "Tracks sowing dates, seed varieties, germination rates, and expected harvest targets."),
        ("Field Activity Records", "Maintains chronological logs of tillage, sowing, spraying, weeding, and fertigation."),
        ("Lifecycle Milestone Tracking", "Visual progress bar of crop phenology stages from emergence to physiological maturity."),
        ("Automated Reminders", "Proactive alerts for upcoming spray schedules, water turns, and inspection routines."),
        ("Historical Farm Archive", "Past season yields, input expenses, and profitability records for continuous learning.")
    ]

    for idx, (ftit, fdesc) in enumerate(fm_features):
        r = idx // 3
        c = idx % 3
        fx = Inches(0.8) + c * (Inches(3.75) + Inches(0.24))
        fy = Inches(1.85) + r * (Inches(2.5) + Inches(0.2))

        add_card(s17, fx, fy, Inches(3.75), Inches(2.4),
                 title=ftit,
                 body_items=[fdesc],
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 18: DAILY FIELD ACTIVITY
    # =========================================================================
    s18 = prs.slides.add_slide(blank_layout)
    add_slide_background(s18)
    add_header(s18, 18, "Field Execution", "Daily Field Activity Checklist", "A dynamic interactive task manager keeping field teams synchronized.")

    add_card(s18, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Today's Active Farm Tasks",
             body_items=[
                 "[✓] 06:00 AM — Run Zone A drip irrigation (Completed)",
                 "[✓] 07:30 AM — Visual scouting for leaf hopper insects (Completed)",
                 "[➔] 10:00 AM — Foliar Micronutrient spray on Block 2 (In Progress)",
                 "[ ] 03:00 PM — Soil moisture sensor data sync (Scheduled)",
                 "[ ] 05:30 PM — Tractor diesel check & air filter clean (Pending)",
                 "— Ensures no critical agricultural intervention is missed or delayed."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_GREEN)

    add_card(s18, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="Checklist Capabilities",
             body_items=[
                 "Role-Based Delegation: Assign tasks to farm workers with SMS/push notification.",
                 "Time-Stamped Completion: Verifies when and where farm inputs were applied.",
                 "Progress Analytics: Visual percentage completion bar for daily farm goals.",
                 "Smart Priority Sorting: Urgency dynamically adjusted by real-time weather forecasts.",
                 "Audit Trail: Essential for organic farming and export quality certification."
             ],
             border_color=C_BORDER,
             bg_color=C_CARD_ALT,
             title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 19: EQUIPMENT & MACHINERY
    # =========================================================================
    s19 = prs.slides.add_slide(blank_layout)
    add_slide_background(s19)
    add_header(s19, 19, "Asset Management", "Equipment & Machinery Management", "Equipment ➔ Usage ➔ Maintenance ➔ Proactive Alerts")

    eq_steps = [
        ("Equipment Digital Registry", [
            "Tractors, tillers, boom sprayers, pumps & drip filters",
            "Model, purchase year, horsepower & warranty status",
            "Tool attachments and implements inventory"
        ]),
        ("Usage & Telematics Tracking", [
            "Engine operational hours tracking",
            "Fuel consumption and diesel efficiency monitoring",
            "Seasonal wear-and-tear calculations"
        ]),
        ("Maintenance Schedules & Alerts", [
            "Proactive oil change and grease point reminders",
            "Pre-season overhaul alerts before harvest peaks",
            "Avoids costly breakdown delays during tight sowing windows"
        ])
    ]

    for idx, (etit, eitems) in enumerate(eq_steps):
        ex = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s19, ex, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=etit,
                 body_items=eitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 20: FARM-TO-MARKET SUPPLY CHAIN
    # =========================================================================
    s20 = prs.slides.add_slide(blank_layout)
    add_slide_background(s20)
    add_header(s20, 20, "Supply Chain", "Connecting Farmers to Markets", "Farmer ➔ Crop Production ➔ Quality ➔ Market ➔ Buyer ➔ Fair Sale")

    sc_cards = [
        ("Production & Harvest", [
            "Harvest volume estimation",
            "Batch traceability from field",
            "Post-harvest sorting protocols"
        ]),
        ("Quality Verification", [
            "Grade A/B/C standardization",
            "Moisture percentage validation",
            "Pesticide residue safety check"
        ]),
        ("Direct Buyer Matching", [
            "Direct connect with food processors",
            "Wholesale aggregators & retail chains",
            "Eliminates 3-4 predatory middlemen layers"
        ]),
        ("Transparent Logistics", [
            "Cold chain & transport coordination",
            "Live transit tracking to delivery hub",
            "Guaranteed digital escrow payouts"
        ])
    ]

    sc_w = Inches(2.78)
    for idx, (sctit, scitems) in enumerate(sc_cards):
        scx = Inches(0.8) + idx * (sc_w + Inches(0.18))
        add_card(s20, scx, Inches(1.85), sc_w, Inches(5.1),
                 title=sctit,
                 body_items=scitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 21: AGRICULTURAL MARKETPLACE
    # =========================================================================
    s21 = prs.slides.add_slide(blank_layout)
    add_slide_background(s21)
    add_header(s21, 21, "Commerce", "Digital Agricultural Marketplace", "A two-way unified trade network for farm produce and agricultural inputs.")

    add_card(s21, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Sell Side: Farm Produce",
             body_items=[
                 "Direct Produce Listings: Farmers post harvested grains, fruits, and vegetables.",
                 "Price Discovery: Market-benchmarked pricing prevents distress selling.",
                 "Bulk Procurement: Connects agricultural cooperatives with institutional buyers.",
                 "Escrow Payment Protection: Ensures farmers are paid immediately upon delivery."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_GREEN)

    add_card(s21, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="Buy Side: Agricultural Inputs & Services",
             body_items=[
                 "Certified Seeds & Fertilizers: Verified authentic inputs directly from suppliers.",
                 "Rental Equipment: Access heavy harvesters and drones on pay-per-acre models.",
                 "Agronomy Services: On-demand soil testing and drone spraying services.",
                 "Transparent Reviews: Verified farmer ratings on agricultural suppliers."
             ],
             border_color=C_ACCENT_CYAN,
             bg_color=C_CARD_ALT,
             title_color=C_ACCENT_CYAN)

    # =========================================================================
    # SLIDE 22: LIVESTOCK & DAIRY
    # =========================================================================
    s22 = prs.slides.add_slide(blank_layout)
    add_slide_background(s22)
    add_header(s22, 22, "Animal Husbandry", "Beyond Crop Farming", "Comprehensive digital management for dairy cattle, livestock health, and milk yields.")

    ls_cards = [
        ("Animal Profiles & RFID Tagging", [
            "Individual cattle records & breed identification",
            "Age, lactation cycle, and pedigree history",
            "Digital tagging of dairy herds"
        ]),
        ("Feed & Nutritional Optimization", [
            "Daily green and dry fodder intake planning",
            "Concentrate and mineral mixture schedules",
            "Nutritional balancing to maximize milk production"
        ]),
        ("Health & Productivity Monitoring", [
            "Vaccination & deworming calendar alerts",
            "Artificial insemination (AI) heat cycle tracking",
            "Daily milk yield recording & mastitis detection"
        ])
    ]

    for idx, (lstit, lsitems) in enumerate(ls_cards):
        lsx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s22, lsx, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=lstit,
                 body_items=lsitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 23: SATELLITE & DRONE INTELLIGENCE
    # =========================================================================
    s23 = prs.slides.add_slide(blank_layout)
    add_slide_background(s23)
    add_header(s23, 23, "Remote Sensing", "Farm Intelligence From Above", "Satellite/Drone Data ➔ AI Analysis ➔ Field Insights ➔ Farmer Action")

    rs_cards = [
        ("Multispectral Satellite Imagery", [
            "Sentinel & Landsat vegetation indices",
            "NDVI (Normalized Difference Vegetation Index) maps",
            "NDRE & EVI for deep canopy assessment"
        ]),
        ("High-Res Drone Orthomosaics", [
            "Centimeter-level aerial field scouting",
            "Early weed patch identification",
            "Crop lodging and flood damage inspection"
        ]),
        ("Macro-to-Micro Actionable Insights", [
            "Identifies hidden nutrient deficiency zones",
            "Detects irrigation pipe leaks and dry spots",
            "Directs ground scouting only where needed"
        ])
    ]

    for idx, (rstit, rsitems) in enumerate(rs_cards):
        rsx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s23, rsx, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=rstit,
                 body_items=rsitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 24: SUSTAINABLE FARMING
    # =========================================================================
    s24 = prs.slides.add_slide(blank_layout)
    add_slide_background(s24)
    add_header(s24, 24, "Regenerative Agronomy", "Smarter Farming. Sustainable Future.", "Data-driven stewardship conserving resources for the next generation.")

    sust_cards = [
        ("Water Stewardship", [
            "Precision drip timing stops aquifer depletion",
            "Zero over-irrigation preserves soil aeration",
            "Rainwater harvesting guidance"
        ]),
        ("Reduced Chemical Inputs", [
            "Targeted spot-spraying replaces blanket pesticide bombing",
            "Calculated NPK limits prevent nitrate runoff",
            "Biological & neem-based pest control prioritization"
        ]),
        ("Soil Regeneration", [
            "Cover cropping and crop rotation recommendations",
            "Organic matter & humus restoration pathways",
            "Long-term soil microbiome health tracking"
        ])
    ]

    for idx, (stit, sitems) in enumerate(sust_cards):
        sx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s24, sx, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=stit,
                 body_items=sitems,
                 border_color=C_ACCENT_GREEN if idx == 1 else C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 25: FINANCE & INSURANCE
    # =========================================================================
    s25 = prs.slides.add_slide(blank_layout)
    add_slide_background(s25)
    add_header(s25, 25, "Financial Intelligence", "Agricultural Financial Intelligence", "De-risking farm operations through institutional credit readiness and insurance awareness.")

    fin_cards = [
        ("Credit Score & Kisan Credit (KCC)", [
            "Digital farm activity ledger proves creditworthiness",
            "Simplifies institutional bank loan paperwork",
            "Reduces reliance on exploitative local moneylenders"
        ]),
        ("Crop Insurance (PMFBY / Micro-insurance)", [
            "Educates farmers on claim deadlines and crop coverage",
            "Geo-tagged weather and satellite proof for fast claims",
            "Simplifies drought and hail damage assessment"
        ]),
        ("Input Cost & Margin Forecasting", [
            "Per-acre cost of cultivation calculator",
            "Tracks operational expenses vs projected revenues",
            "Supports strategic business decisions for farm families"
        ])
    ]

    for idx, (ftit, fitems) in enumerate(fin_cards):
        fx = Inches(0.8) + idx * (Inches(3.75) + Inches(0.24))
        add_card(s25, fx, Inches(1.85), Inches(3.75), Inches(5.1),
                 title=ftit,
                 body_items=fitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 26: ANALYTICS & INSIGHTS
    # =========================================================================
    s26 = prs.slides.add_slide(blank_layout)
    add_slide_background(s26)
    add_header(s26, 26, "Data Intelligence", "Data Into Decisions", "Turning farm telemetry and historical records into clear visual business intelligence.")

    analytics_grid = [
        ("Farm Yield Trends", ["Year-on-year crop productivity comparison", "Benchmarking against regional taluk averages", "Variety performance analysis"]),
        ("Resource Efficiency", ["Liters of irrigation applied per kilogram produced", "Fertilizer utilization efficiency percentage", "Diesel & labor cost breakdown"]),
        ("Market Price Volatility", ["Historical mandi pricing curves", "Optimal selling month identification", "Realized margin tracking"]),
        ("AI Diagnostic Patterns", ["Prevalent seasonal pest outbreak timelines", "Soil pH shift trajectory across 3 years", "Predictive maintenance schedule accuracy"])
    ]

    for idx, (atit, aitems) in enumerate(analytics_grid):
        ax = Inches(0.8) + (idx % 2) * (Inches(5.7) + Inches(0.333))
        ay = Inches(1.85) + (idx // 2) * (Inches(2.5) + Inches(0.2))

        add_card(s26, ax, ay, Inches(5.7), Inches(2.45),
                 title=atit,
                 body_items=aitems,
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 27: ALERTS & NOTIFICATIONS
    # =========================================================================
    s27 = prs.slides.add_slide(blank_layout)
    add_slide_background(s27)
    add_header(s27, 27, "Early Warning System", "Proactive Agricultural Alerts", "Detect ➔ Analyze ➔ Alert ➔ Act: Preventing farm losses before they happen.")

    alert_steps = [
        ("01. Continuous Detection", [
            "Real-time sensor threshold triggers",
            "Satellite anomaly detection feeds",
            "Sudden weather radar storm alerts",
            "Mandi market price spike/dip triggers"
        ]),
        ("02. Contextual AI Analysis", [
            "Evaluates impact specifically on farmer's crop",
            "Filters noise to avoid notification fatigue",
            "Calculates time window available for farmer action"
        ]),
        ("03. Multi-Channel Alert Delivery", [
            "High-priority push notifications",
            "Native audio voice alerts in regional languages",
            "SMS alerts for low-connectivity offline areas"
        ]),
        ("04. Guided Immediate Action", [
            "Exact step-by-step mitigation instructions",
            "Direct link to relevant equipment/spray checklist",
            "Follow-up verification check scheduled"
        ])
    ]

    aw_w = Inches(2.78)
    for idx, (atit, aitems) in enumerate(alert_steps):
        ax = Inches(0.8) + idx * (aw_w + Inches(0.18))
        add_card(s27, ax, Inches(1.85), aw_w, Inches(5.1),
                 title=atit,
                 body_items=aitems,
                 border_color=C_ACCENT_GREEN if idx == 3 else C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 28: COMPLETE USER WORKFLOW
    # =========================================================================
    s28 = prs.slides.add_slide(blank_layout)
    add_slide_background(s28)
    add_header(s28, 28, "User Experience", "How AgriOne Works – Complete User Workflow", "An intuitive 11-step journey connecting login to tangible farm action.")

    flow_steps = [
        "1. Farmer Login", "2. Language Selection", "3. Farm Location", "4. Farm Profile",
        "5. Dashboard Overview", "6. Data Aggregation", "7. AI Processing", "8. Rule Validation",
        "9. Actionable Advice", "10. Farmer Action", "11. Data Feedback Loop"
    ]

    f_box_w = Inches(1.85)
    f_gap = Inches(0.12)
    for i, step in enumerate(flow_steps[:6]):
        fx = Inches(0.8) + i * (f_box_w + f_gap)
        card = s28.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, Inches(2.1), f_box_w, Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_BG
        card.line.color.rgb = C_ACCENT_CYAN
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = step
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_TEXT_WHITE

    for i, step in enumerate(flow_steps[6:]):
        fx = Inches(0.8) + i * (Inches(2.25) + Inches(0.12))
        card = s28.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, Inches(4.3), Inches(2.25), Inches(1.8))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD_ALT
        card.line.color.rgb = C_ACCENT_GREEN
        card.line.width = Pt(1)
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = step
        p.font.name = FONT_TITLE
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = C_HIGHLIGHT

    # =========================================================================
    # SLIDE 29: TECHNOLOGY ARCHITECTURE
    # =========================================================================
    s29 = prs.slides.add_slide(blank_layout)
    add_slide_background(s29)
    add_header(s29, 29, "System Architecture", "AgriOne Technology Stack", "A robust, scalable, enterprise-grade layered technical architecture.")

    arch_layers = [
        ("Layer 1: Frontend UI/UX", "Responsive Web App • Multilingual UI • Voice I/O Engine • Telemetry Dashboard", C_ACCENT_CYAN),
        ("Layer 2: Backend & APIs", "Python Backend Architecture • FastAPI / REST Services • JWT Auth • Task Queue Processing", C_TEXT_WHITE),
        ("Layer 3: AI Intelligence Layer", "Computer Vision Disease Classifier • LLM Agronomy Assistant • Multilingual ASR & TTS", C_HIGHLIGHT),
        ("Layer 4: Agronomic Rule Engine", "Calculated Soil Equations • Water Infiltration Rules • Chemical Dosage Guardrails", C_ACCENT_GREEN),
        ("Layer 5: Unified Data Layer", "Farmer Profiles • Sensor Telemetry • Geospatial Polygons • Mandi Price Time-Series", C_TEXT_WHITE),
        ("Layer 6: External Integrations", "Meteorological Weather APIs • Satellite NDVI Feeds • Mandi APMC APIs • Location Services", C_ACCENT_CYAN)
    ]

    layer_h = Inches(0.72)
    layer_gap = Inches(0.12)
    for idx, (ltit, ldesc, lcolor) in enumerate(arch_layers):
        ly = Inches(1.85) + idx * (layer_h + layer_gap)
        lcard = s29.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), ly, Inches(11.733), layer_h)
        lcard.fill.solid()
        lcard.fill.fore_color.rgb = C_CARD_BG
        lcard.line.color.rgb = lcolor
        lcard.line.width = Pt(1)

        tf = lcard.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        p1 = tf.paragraphs[0]
        p1.text = ltit
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = lcolor

        p2 = tf.add_paragraph()
        p2.text = ldesc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.space_before = Pt(1)

    # =========================================================================
    # SLIDE 30: AI + RULE ENGINE
    # =========================================================================
    s30 = prs.slides.add_slide(blank_layout)
    add_slide_background(s30)
    add_header(s30, 30, "Technical Core", "Why AgriOne Is More Than a Chatbot", "The critical distinction between generic AI talk and validated agronomic decision support.")

    # Top comparison banner
    add_card(s30, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="The Role of the AI Layer",
             body_items=[
                 "Understands natural, unstructured farmer language across multiple dialects.",
                 "Interprets messy camera images of infected crop leaves.",
                 "Extracts farmer intent: crop variety, symptom description, field stage.",
                 "Synthesizes conversational voice and regional text replies.",
                 "Limitation alone: LLMs can hallucinate doses or ignore regional soil laws."
             ],
             border_color=C_ACCENT_CYAN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_CYAN)

    add_card(s30, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="The Role of the Rule Engine",
             body_items=[
                 "Validates agronomic rules: checks maximum chemical toxicity thresholds.",
                 "Executes exact mathematical calculations for irrigation liters and fertilizer kg/acre.",
                 "Enforces agro-climatic boundaries: never suggests frost crops in tropical heat.",
                 "Grounds the AI: Prevents harmful advice, ensuring 100% safe farm operations.",
                 "Combined Result: AgriOne provides reliable, verifiable decision support."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_ALT,
             title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 31: KEY DIFFERENTIATORS
    # =========================================================================
    s31 = prs.slides.add_slide(blank_layout)
    add_slide_background(s31)
    add_header(s31, 31, "Competitive Edge", "Why AgriOne?", "“AgriOne converts fragmented agricultural data into practical farming decisions.”")

    diffs = [
        ("Unified Ecosystem", "Consolidates 18 agricultural tools into one cohesive platform instead of 10 isolated apps."),
        ("AI + Rule Engine Fusion", "Marries conversational NLP with strict mathematical agronomy for 100% verified safety."),
        ("Voice-First & Multilingual", "Fully usable by farmers of all literacy levels in their native regional language."),
        ("Camera-Based Diagnostics", "Live leaf image scanning for rapid, laboratory-grade disease detection."),
        ("Hyper-Local Weather", "Transforms abstract weather data into exact irrigation and spraying schedules."),
        ("Complete Farm Lifecycle", "Guides farmers from soil preparation to mandi sale and livestock management.")
    ]

    for idx, (dtit, ddesc) in enumerate(diffs):
        r = idx // 3
        c = idx % 3
        dx = Inches(0.8) + c * (Inches(3.75) + Inches(0.24))
        dy = Inches(1.85) + r * (Inches(2.5) + Inches(0.2))

        add_card(s31, dx, dy, Inches(3.75), Inches(2.4),
                 title=dtit,
                 body_items=[ddesc],
                 border_color=C_BORDER,
                 bg_color=C_CARD_BG,
                 title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 32: BENEFITS & IMPACT
    # =========================================================================
    s32 = prs.slides.add_slide(blank_layout)
    add_slide_background(s32)
    add_header(s32, 32, "Value Creation", "Impact of AgriOne", "Delivering measurable transformation across individual livelihoods and national agriculture.")

    add_card(s32, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Tangible Impact For Farmers",
             body_items=[
                 "Better Decisions: Science-backed crop selection and precise fertilization.",
                 "Reduced Uncertainty: Eliminates guesswork around weather and pest damage.",
                 "Cost Reduction: Cuts fertilizer waste and diesel consumption significantly.",
                 "Market Transparency: Prevents middleman exploitation, increasing net income.",
                 "Accessibility: Voice interaction in native tongue empowers all farmers."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_BG,
             title_color=C_ACCENT_GREEN)

    add_card(s32, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="Systemic Impact For Agriculture",
             body_items=[
                 "Water Conservation: Protects depleting aquifers through smart drip schedules.",
                 "Soil Preservation: Mitigates soil salinization and toxic chemical runoff.",
                 "Data-Driven Agronomy: Accelerates adoption of modern digital techniques.",
                 "Enhanced Food Security: Reduces post-harvest losses and disease crop failure.",
                 "Sustainable Future: Builds resilient farms adaptable to changing climates."
             ],
             border_color=C_ACCENT_CYAN,
             bg_color=C_CARD_ALT,
             title_color=C_ACCENT_CYAN)

    # =========================================================================
    # SLIDE 33: TRADITIONAL FARMING VS AGRIONE
    # =========================================================================
    s33 = prs.slides.add_slide(blank_layout)
    add_slide_background(s33)
    add_header(s33, 33, "Comparative Analysis", "Traditional Farming vs AgriOne", "The fundamental shift from fragmented intuition to data-guided precision.")

    add_card(s33, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1),
             title="Traditional Farming (The Past)",
             body_items=[
                 "Fragmented Information: Scattered across gossip, TV news, and guesswork.",
                 "Manual Paper Tracking: Lost spray dates, forgotten water turns, no records.",
                 "Delayed Decisions: Pest damage only treated after half the field is ruined.",
                 "Opaque Market Prices: Forced to accept lowest prices offered by local buyers.",
                 "High Language Barrier: Technical research inaccessible to regional speakers.",
                 "Reactive Mindset: Constantly scrambling to respond to farm crises."
             ],
             border_color=C_BORDER_ALERT,
             bg_color=C_CARD_ALERT,
             title_color=C_ACCENT_RED)

    add_card(s33, Inches(6.8), Inches(1.8), Inches(5.733), Inches(5.1),
             title="AgriOne Platform (The Future)",
             body_items=[
                 "Unified Intelligence: Weather, soil, crops, and markets in one command hub.",
                 "Automated Digital Ledger: Real-time checklists and historical telemetry.",
                 "Instant AI Diagnostics: Leaf image scan catches infections within seconds.",
                 "Live Mandi Comparison: Benchmarks multiple APMC prices for highest profit.",
                 "Native Voice & Multilingual: Fluid interaction in farmer's everyday language.",
                 "Proactive Predictive Approach: Early warnings avert damage before it occurs."
             ],
             border_color=C_ACCENT_GREEN,
             bg_color=C_CARD_ALT,
             title_color=C_HIGHLIGHT)

    # =========================================================================
    # SLIDE 34: FUTURE SCOPE
    # =========================================================================
    s34 = prs.slides.add_slide(blank_layout)
    add_slide_background(s34)
    add_header(s34, 34, "Roadmap", "Future of AgriOne", "Clearly defined Future Scope – Planned extensions for next project phases.")

    future_pillars = [
        ("IoT & Autonomous Automation", [
            "Hardware soil probe integration (LoRaWAN & NB-IoT)",
            "Automated solenoid irrigation valve control",
            "Solar-powered weather station telemetry"
        ]),
        ("Advanced Aerial Analytics", [
            "Autonomous drone swarm crop spraying",
            "High-frequency synthetic aperture radar (SAR) satellite data",
            "Hyper-spectral plant stress detection"
        ]),
        ("Predictive AI Models", [
            "Machine-learning yield forecasting per acre",
            "Deep learning multi-week mandi price forecasting",
            "Regional pest infestation epidemic tracking"
        ]),
        ("Policy & Ecosystem Extension", [
            "Direct government subsidy & scheme integration",
            "Decentralized peer-to-peer farmer community forums",
            "Automated smart-contract micro-insurance payouts"
        ])
    ]

    fw_w = Inches(2.78)
    for idx, (ftit, fitems) in enumerate(future_pillars):
        fx = Inches(0.8) + idx * (fw_w + Inches(0.18))
        add_card(s34, fx, Inches(1.85), fw_w, Inches(5.1),
                 title=ftit,
                 body_items=fitems,
                 border_color=C_ACCENT_CYAN,
                 bg_color=C_CARD_BG,
                 title_color=C_ACCENT_CYAN)

    # =========================================================================
    # SLIDE 35: CONCLUSION
    # =========================================================================
    s35 = prs.slides.add_slide(blank_layout)
    add_slide_background(s35)

    # Main Conclusion Card
    c_card = s35.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), Inches(11.733), Inches(5.5))
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = C_CARD_ALT
    c_card.line.color.rgb = C_ACCENT_GREEN
    c_card.line.width = Pt(2)

    tf_c = c_card.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = Inches(0.5)
    tf_c.margin_top = Inches(0.5)

    p_ctag = tf_c.paragraphs[0]
    p_ctag.text = "AGRIONE • CONCLUSION & VISION"
    p_ctag.font.name = FONT_TITLE
    p_ctag.font.size = Pt(13)
    p_ctag.font.bold = True
    p_ctag.font.color.rgb = C_ACCENT_GREEN

    p_ctit = tf_c.add_paragraph()
    p_ctit.text = "One Platform.  Multiple Agricultural Intelligence Systems.  One Smarter Farming Ecosystem."
    p_ctit.font.name = FONT_TITLE
    p_ctit.font.size = Pt(24)
    p_ctit.font.bold = True
    p_ctit.font.color.rgb = C_TEXT_WHITE
    p_ctit.space_before = Pt(14)
    p_ctit.space_after = Pt(14)

    p_cmsg = tf_c.add_paragraph()
    p_cmsg.text = "“AgriOne aims to empower farmers with accessible, intelligent and actionable agricultural technology — transforming raw data into better decisions and better decisions into a sustainable, prosperous agricultural future.”"
    p_cmsg.font.name = FONT_BODY
    p_cmsg.font.size = Pt(16)
    p_cmsg.font.italic = True
    p_cmsg.font.color.rgb = C_ACCENT_CYAN
    p_cmsg.space_after = Pt(20)

    p_sum = tf_c.add_paragraph()
    p_sum.text = "By bridging multimodal AI with rule-based agronomy, voice accessibility, and market transparency, AgriOne stands as a comprehensive digital companion for modern agriculture."
    p_sum.font.name = FONT_BODY
    p_sum.font.size = Pt(13)
    p_sum.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 36: THANK YOU
    # =========================================================================
    s36 = prs.slides.add_slide(blank_layout)
    add_slide_background(s36)

    # Thank You Banner & Team
    ty_card = s36.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    ty_card.fill.solid()
    ty_card.fill.fore_color.rgb = C_CARD_BG
    ty_card.line.color.rgb = C_ACCENT_GREEN
    ty_card.line.width = Pt(1.5)

    tf_ty = ty_card.text_frame
    tf_ty.word_wrap = True
    tf_ty.margin_left = Inches(0.5)
    tf_ty.margin_top = Inches(0.4)

    p_ty1 = tf_ty.paragraphs[0]
    p_ty1.text = "THANK YOU"
    p_ty1.font.name = FONT_TITLE
    p_ty1.font.size = Pt(40)
    p_ty1.font.bold = True
    p_ty1.font.color.rgb = C_TEXT_WHITE

    p_ty2 = tf_ty.add_paragraph()
    p_ty2.text = "AgriOne – AI-Powered Intelligence for Smarter Farming."
    p_ty2.font.name = FONT_TITLE
    p_ty2.font.size = Pt(18)
    p_ty2.font.bold = True
    p_ty2.font.color.rgb = C_ACCENT_CYAN
    p_ty2.space_before = Pt(4)
    p_ty2.space_after = Pt(16)

    p_ty3 = tf_ty.add_paragraph()
    p_ty3.text = "TEAM: SILENT CODERS"
    p_ty3.font.name = FONT_TITLE
    p_ty3.font.size = Pt(14)
    p_ty3.font.bold = True
    p_ty3.font.color.rgb = C_ACCENT_GREEN

    p_lead = tf_ty.add_paragraph()
    p_lead.text = "• Iyaan Nagari  (Team Leader)"
    p_lead.font.name = FONT_TITLE
    p_lead.font.size = Pt(14)
    p_lead.font.bold = True
    p_lead.font.color.rgb = C_TEXT_WHITE
    p_lead.space_before = Pt(6)

    for m in ["Bhuvan Kadakabhavi", "Jafar Momin", "Aiman Nadaf"]:
        p_mem = tf_ty.add_paragraph()
        p_mem.text = f"• {m}  (Team Member)"
        p_mem.font.name = FONT_BODY
        p_mem.font.size = Pt(13)
        p_mem.font.color.rgb = C_TEXT_MUTED
        p_mem.space_before = Pt(3)

    p_open = tf_ty.add_paragraph()
    p_open.text = "\nQuestions & Technical Discussion Welcomed."
    p_open.font.name = FONT_TITLE
    p_open.font.size = Pt(13)
    p_open.font.bold = True
    p_open.font.color.rgb = C_HIGHLIGHT

    # Add Silent Coders Logo on Right of Slide 36
    if os.path.exists(sc_logo_path):
        s36.shapes.add_picture(sc_logo_path, Inches(8.5), Inches(1.6), Inches(3.6), Inches(3.6))

    # Save presentation
    output_path = os.path.join(workspace_dir, "AgriOne_Silent_Coders_Presentation.pptx")
    prs.save(output_path)
    print(f"Successfully generated {output_path} with {len(prs.slides)} slides.")

if __name__ == "__main__":
    create_deck()
